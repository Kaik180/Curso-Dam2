import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    # Formato panorámico moderno 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colores sobrios y limpios
    COLOR_PRIMARY = RGBColor(20, 50, 90)     # Azul marino oscuro
    COLOR_TEXT = RGBColor(40, 40, 40)        # Gris oscuro para lectura cómoda
    COLOR_MUTED = RGBColor(100, 100, 100)    # Gris secundario
    COLOR_ACCENT = RGBColor(40, 80, 130)     # Azul medio

    blank_layout = prs.slide_layouts[6] # Diseño en blanco

    current_dir = os.path.dirname(os.path.abspath(__file__))
    logo_path = os.path.join(current_dir, "sabor_urbano_logo.jpg")

    def add_header(slide, title_text, category_text="DIGITALIZACIÓN APLICADA AL SECTOR PRODUCTIVO"):
        # Categoría pequeña arriba
        tx_box_cat = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(0.4))
        tf_cat = tx_box_cat.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        # Título principal de la diapositiva
        tx_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(11.3), Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY

    # ==========================================
    # DIAPOSITIVA 1: PORTADA
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Logo a la izquierda
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(1.2), Inches(1.8), width=Inches(3.8))

    # Textos a la derecha
    tx_box1 = slide1.shapes.add_textbox(Inches(5.5), Inches(1.8), Inches(6.8), Inches(4.2))
    tf1 = tx_box1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "PROYECTO DE DIGITALIZACIÓN DE EMPRESA"
    p0.font.name = "Arial"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "Integración de Tecnologías IT y OT:\nRestaurante \"Sabor Urbano\""
    p1.font.name = "Arial"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY
    p1.space_after = Pt(24)

    p2 = tf1.add_paragraph()
    p2.text = "Alumno: Kevin\nCurso: 2º DAM (Desarrollo de Aplicaciones Multiplataforma)\nMódulo: Digitalización Aplicada al Sector Productivo\nCurso Académico: 2025 - 2026"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_TEXT

    # ==========================================
    # DIAPOSITIVA 2: PUNTO DE PARTIDA Y PROBLEMAS
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Situación Inicial y Problemas Detectados")

    tx2 = slide2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
    tf2 = tx2.text_frame
    tf2.word_wrap = True

    points2 = [
        ("Restaurante tradicional bien situado:", "Ubicado en el centro de la ciudad con buena comida mediterránea y de fusión, pero gestionado totalmente a mano."),
        ("Reservas en libreta de papel:", "Anotadas a mano o por teléfono; causaba mesas duplicadas, despistes y huecos vacíos por cancelaciones sin aviso."),
        ("Control visual del almacén:", "Se revisaba a ojo sin apuntar nada; se caducaba comida, faltaba género en horas punta y no se controlaban las mermas."),
        ("Caja registradora antigua:", "Cobros aislados sin conexión a ningún programa ni contabilidad, con cierres lentos y descuadres de dinero."),
        ("Comandas en papel y a voz:", "Notas a mano ilegibles que se perdían, errores con platos para alérgicos y retrasos entre la sala y la cocina.")
    ]

    for i, (title, desc) in enumerate(points2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_after = Pt(14)
        run_b = p.add_run()
        run_b.text = "•  " + title + " "
        run_b.font.name = "Arial"
        run_b.font.bold = True
        run_b.font.size = Pt(15)
        run_b.font.color.rgb = COLOR_PRIMARY

        run_t = p.add_run()
        run_t.text = desc
        run_t.font.name = "Arial"
        run_t.font.size = Pt(14)
        run_t.font.color.rgb = COLOR_TEXT

    # ==========================================
    # DIAPOSITIVA 3: TECNOLOGÍAS OT (EQUIPOS FÍSICOS)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. Equipos Físicos en el Local (Tecnologías OT)")

    tx3 = slide3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
    tf3 = tx3.text_frame
    tf3.word_wrap = True

    points3 = [
        ("Comanderos táctiles (Sunmi V2s Plus):", "3 terminales móviles para camareros. Toman la comanda en la mesa y llega al segundo a la cocina."),
        ("Pantallas táctiles de cocina KDS (ELO Touch 22\"): ", "2 pantallas resistentes a grasa y vapor en las partidas de cocina. Organizan los platos por orden y tiempo."),
        ("Cajón de cobro automático (Cashlogy POS 1500):", "Cobra el efectivo, devuelve el cambio exacto, rechaza billetes falsos y cuadra la caja al céntimo."),
        ("Control de almacén (Zebra TC21 + Báscula Dibal):", "Lectores de códigos de barras y báscula conectada para registrar pedidos de proveedores y pesar materias primas."),
        ("Sensores de frío para cámaras (Testo Saveris):", "Monitorizan la temperatura 24/7 y avisan al móvil si se corta la luz para cumplir la normativa sanitaria APPCC.")
    ]

    for i, (title, desc) in enumerate(points3):
        p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        p.space_after = Pt(14)
        run_b = p.add_run()
        run_b.text = "•  " + title + " "
        run_b.font.name = "Arial"
        run_b.font.bold = True
        run_b.font.size = Pt(15)
        run_b.font.color.rgb = COLOR_PRIMARY

        run_t = p.add_run()
        run_t.text = desc
        run_t.font.name = "Arial"
        run_t.font.size = Pt(14)
        run_t.font.color.rgb = COLOR_TEXT

    # ==========================================
    # DIAPOSITIVA 4: TECNOLOGÍAS IT (SOFTWARE Y NUBE)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Programas y Servicios en la Nube (Tecnologías IT)")

    tx4 = slide4.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
    tf4 = tx4.text_frame
    tf4.word_wrap = True

    points4 = [
        ("Programa de Gestión ERP (Revo Xef + Holded):", "Control de compras, facturas con el móvil por OCR, cálculo del coste real de cada plato (escandallos) y factura electrónica legal."),
        ("Gestión de Reservas Online (CoverManager):", "Reservas 24 horas desde la web, Google Maps o Instagram sin pagar comisiones; organiza las mesas y pide fianza para evitar mesas vacías."),
        ("Página Web propia con pedidos (WordPress + WooCommerce):", "Carta digital actualizada con fotos y venta directa de comida para llevar sin intermediarios."),
        ("Presencia local y opiniones (Localoo / Google Business):", "Mantiene siempre al día los horarios y fotos en Google Maps y ayuda a responder rápido a las reseñas de los clientes."),
        ("Cuadro de Mandos con Gráficos (Power BI):", "Muestra en pantalla qué platos dejan más dinero limpio, qué días se llena más y cuánta comida comprar para no tirar nada.")
    ]

    for i, (title, desc) in enumerate(points4):
        p = tf4.paragraphs[0] if i == 0 else tf4.add_paragraph()
        p.space_after = Pt(14)
        run_b = p.add_run()
        run_b.text = "•  " + title + " "
        run_b.font.name = "Arial"
        run_b.font.bold = True
        run_b.font.size = Pt(15)
        run_b.font.color.rgb = COLOR_PRIMARY

        run_t = p.add_run()
        run_t.text = desc
        run_t.font.name = "Arial"
        run_t.font.size = Pt(14)
        run_t.font.color.rgb = COLOR_TEXT

    # ==========================================
    # DIAPOSITIVA 5: FASES DE IMPLANTACIÓN
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. Plan de Trabajo en 5 Fases (8 Semanas)")

    tx5 = slide5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
    tf5 = tx5.text_frame
    tf5.word_wrap = True

    points5 = [
        ("Fase 1 (Semanas 1-2) - Red y Comunicaciones:", "Instalar cable de red Cat 6A, dos puntos de acceso Wi-Fi 6 para los comanderos y router con SIM 4G/5G de respaldo."),
        ("Fase 2 (Semanas 3-4) - Configuración del Software:", "Dar de alta el ERP, meter la carta con precios, ingredientes y alérgenos, y configurar el mapa de mesas."),
        ("Fase 3 (Semanas 5-6) - Instalación de Aparatos:", "Montar las pantallas en cocina, el cajón automático en barra, configurar los comanderos y colocar los sensores de frío."),
        ("Fase 4 (Semana 7) - Pruebas y Ensayo:", "Comprobar que todo sincroniza y hacer un servicio a puerta cerrada con familiares para probar el sistema en real."),
        ("Fase 5 (Semana 8) - Formación y Puesta en Marcha:", "Enseñar a usar las aplicaciones a camareros y cocineros, abrir con el sistema nuevo y soporte técnico las dos primeras semanas.")
    ]

    for i, (title, desc) in enumerate(points5):
        p = tf5.paragraphs[0] if i == 0 else tf5.add_paragraph()
        p.space_after = Pt(14)
        run_b = p.add_run()
        run_b.text = "•  " + title + " "
        run_b.font.name = "Arial"
        run_b.font.bold = True
        run_b.font.size = Pt(15)
        run_b.font.color.rgb = COLOR_PRIMARY

        run_t = p.add_run()
        run_t.text = desc
        run_t.font.name = "Arial"
        run_t.font.size = Pt(14)
        run_t.font.color.rgb = COLOR_TEXT

    # ==========================================
    # DIAPOSITIVA 6: SEGURIDAD Y PLAN ANTE FALLOS
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. Ciberseguridad y Prevención de Problemas")

    tx6 = slide6.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
    tf6 = tx6.text_frame
    tf6.word_wrap = True

    points6 = [
        ("Miedo al cambio de los trabajadores:", "Se resuelve con formación práctica paso a paso y pantallas con fotos grandes y botones sencillos."),
        ("Si se corta internet o la luz:", "Un SAI mantiene encendida la caja y el router, la conexión cambia sola a 4G/5G en segundos y el TPV puede seguir cobrando en red local sin internet."),
        ("Redes separadas en 4 VLANs:", "Una red para cobrar y pedir, otra para sensores de cocina, otra para oficina y otra para la Wi-Fi de clientes, para que nadie pueda acceder a los datos."),
        ("Pagos seguros (PCI-DSS):", "Los datáfonos van cifrados y el restaurante nunca almacena números de tarjeta bancaria."),
        ("Copias de seguridad (Regla 3-2-1):", "3 copias de los datos de ventas y clientes, en 2 sitios distintos (disco duro local y nube), guardando 1 copia fuera del local.")
    ]

    for i, (title, desc) in enumerate(points6):
        p = tf6.paragraphs[0] if i == 0 else tf6.add_paragraph()
        p.space_after = Pt(14)
        run_b = p.add_run()
        run_b.text = "•  " + title + " "
        run_b.font.name = "Arial"
        run_b.font.bold = True
        run_b.font.size = Pt(15)
        run_b.font.color.rgb = COLOR_PRIMARY

        run_t = p.add_run()
        run_t.text = desc
        run_t.font.name = "Arial"
        run_t.font.size = Pt(14)
        run_t.font.color.rgb = COLOR_TEXT

    # ==========================================
    # DIAPOSITIVA 7: PRESUPUESTO Y RENTABILIDAD
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. Presupuesto y Retorno de la Inversión")

    tx7 = slide7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
    tf7 = tx7.text_frame
    tf7.word_wrap = True

    points7 = [
        ("Inversión en Equipos Físicos (OT):", "11.665 € (3 comanderos, 2 pantallas de cocina, cajón automático Cashlogy, TPV, lectores Zebra y sensores de frío)."),
        ("Inversión en Software y Servicios (IT):", "1.750 € de instalación inicial + 4.684 € al año en programas (ERP, reservas, web y publicidad)."),
        ("Ahorro directo en almacén:", "Se reduce un 15% la comida que se tiraba caducada gracias al control de compras y escandallos."),
        ("Aumento de ventas (+20%):", "Mayor rotación de mesas al mediodía (cobros en 2 minutos) y reservas online 24/7 sin perder clientes por teléfono."),
        ("Recuperación del dinero:", "La inversión total se amortiza en aproximadamente 11 meses de funcionamiento.")
    ]

    for i, (title, desc) in enumerate(points7):
        p = tf7.paragraphs[0] if i == 0 else tf7.add_paragraph()
        p.space_after = Pt(14)
        run_b = p.add_run()
        run_b.text = "•  " + title + " "
        run_b.font.name = "Arial"
        run_b.font.bold = True
        run_b.font.size = Pt(15)
        run_b.font.color.rgb = COLOR_PRIMARY

        run_t = p.add_run()
        run_t.text = desc
        run_t.font.name = "Arial"
        run_t.font.size = Pt(14)
        run_t.font.color.rgb = COLOR_TEXT

    # ==========================================
    # DIAPOSITIVA 8: CONCLUSIONES
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "7. Conclusiones Finales")

    tx8 = slide8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
    tf8 = tx8.text_frame
    tf8.word_wrap = True

    points8 = [
        ("Fin del papel y los gritos:", "Pedidos claros y directos a cocina en segundos, sin confusiones con las alergias."),
        ("Control total del negocio:", "Saber en todo momento qué platos son rentables y cuánta comida comprar con datos reales."),
        ("Clientes mejor atendidos:", "Reservan fácil desde el móvil, los platos salen a su tiempo y el cobro en mesa es inmediato."),
        ("Seguridad garantizada:", "Cumplimiento sanitario estricto en cámaras de frío, dinero en caja protegido y datos a salvo."),
        ("Proyecto viable y realista:", "Una solución práctica pensada para un restaurante real que se paga sola en menos de un año.")
    ]

    for i, (title, desc) in enumerate(points8):
        p = tf8.paragraphs[0] if i == 0 else tf8.add_paragraph()
        p.space_after = Pt(14)
        run_b = p.add_run()
        run_b.text = "•  " + title + " "
        run_b.font.name = "Arial"
        run_b.font.bold = True
        run_b.font.size = Pt(15)
        run_b.font.color.rgb = COLOR_PRIMARY

        run_t = p.add_run()
        run_t.text = desc
        run_t.font.name = "Arial"
        run_t.font.size = Pt(14)
        run_t.font.color.rgb = COLOR_TEXT

    # Guardar en las carpetas
    output_dir = current_dir
    output_file = os.path.join(output_dir, "Presentacion_Sabor_Urbano_Kevin.pptx")
    prs.save(output_file)
    print(f"Presentación guardada en: {output_file}")

    repo_root = os.path.abspath(os.path.join(output_dir, "..", ".."))
    root_copy = os.path.join(repo_root, "Presentacion_Sabor_Urbano_Kevin.pptx")
    folder_copy = os.path.join(repo_root, "digitalizacion-productivo", "Presentacion_Sabor_Urbano_Kevin.pptx")
    
    shutil.copyfile(output_file, root_copy)
    shutil.copyfile(output_file, folder_copy)
    print(f"Copia en la raíz: {root_copy}")
    print(f"Copia en digitalizacion-productivo: {folder_copy}")

if __name__ == "__main__":
    create_presentation()
