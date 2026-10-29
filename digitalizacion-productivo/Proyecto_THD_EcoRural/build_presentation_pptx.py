import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    # Widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Palette
    C_DARK = RGBColor(27, 67, 50)       # #1B4332 Verde bosque profundo
    C_MID = RGBColor(45, 106, 79)       # #2D6A4F Verde medio
    C_ACCENT = RGBColor(64, 145, 108)   # #40916C Verde salvia
    C_TEXT = RGBColor(40, 40, 40)       # Gris texto
    C_MUTED = RGBColor(90, 95, 92)      # Gris secundario
    C_CARD_BG = RGBColor(244, 247, 245) # Fondo tarjeta muy suave
    C_BORDER = RGBColor(200, 215, 206)  # Borde sutil
    C_WHITE = RGBColor(255, 255, 255)
    
    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="DIGITALIZACIÓN APLICADA AL SISTEMA PRODUCTIVO · 2º DAM"):
        # Category label
        tb_cat = slide.shapes.add_textbox(Inches(0.9), Inches(0.5), Inches(11.5), Inches(0.35))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_c = tf_cat.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.name = "Calibri"
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = C_ACCENT
        
        # Main Title
        tb_title = slide.shapes.add_textbox(Inches(0.9), Inches(0.85), Inches(11.5), Inches(0.65))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_t = tf_title.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = C_DARK

    def create_card(slide, left, top, width, height, title, items, subtitle=""):
        # Card background rectangle
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = C_CARD_BG
        shape.line.color.rgb = C_BORDER
        shape.line.width = Pt(1.2)
        
        # Text box inside card
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_head = tf.paragraphs[0]
        p_head.text = title
        p_head.font.name = "Calibri"
        p_head.font.size = Pt(14)
        p_head.font.bold = True
        p_head.font.color.rgb = C_DARK
        p_head.space_after = Pt(2)
        
        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = "Calibri"
            p_sub.font.size = Pt(10)
            p_sub.font.italic = True
            p_sub.font.color.rgb = C_MID
            p_sub.space_after = Pt(8)
        else:
            p_head.space_after = Pt(8)
            
        for bullet in items:
            p_b = tf.add_paragraph()
            p_b.font.name = "Calibri"
            p_b.font.size = Pt(10.5)
            p_b.space_after = Pt(4)
            
            if ":" in bullet:
                b_part, r_part = bullet.split(":", 1)
                run_b = p_b.add_run()
                run_b.text = b_part + ":"
                run_b.font.bold = True
                run_b.font.color.rgb = C_MID
                run_r = p_b.add_run()
                run_r.text = r_part
                run_r.font.color.rgb = C_TEXT
            else:
                run_txt = p_b.add_run()
                run_txt.text = bullet
                run_txt.font.color.rgb = C_TEXT

    # =========================================================================
    # DIAPOSITIVA 1: PORTADA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    
    # Left decorative card
    left_rect = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.2), Inches(4.2), Inches(5.2))
    left_rect.fill.solid()
    left_rect.fill.fore_color.rgb = C_DARK
    left_rect.line.fill.background()
    
    tb_l = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(3.6), Inches(4.0))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    
    p_l1 = tf_l.paragraphs[0]
    p_l1.text = "EcoRural\nConstrucciones"
    p_l1.font.name = "Calibri"
    p_l1.font.size = Pt(32)
    p_l1.font.bold = True
    p_l1.font.color.rgb = C_WHITE
    p_l1.space_after = Pt(18)
    
    p_l2 = tf_l.add_paragraph()
    p_l2.text = "Transformación Digital e Innovación Sostenible en el Medio Rural"
    p_l2.font.name = "Calibri"
    p_l2.font.size = Pt(14)
    p_l2.font.color.rgb = RGBColor(180, 220, 200)
    p_l2.space_after = Pt(24)
    
    p_l3 = tf_l.add_paragraph()
    p_l3.text = "Fundadores: Carlos y Esther\nArquitectura · THD · Smart Village"
    p_l3.font.name = "Calibri"
    p_l3.font.size = Pt(11)
    p_l3.font.color.rgb = C_WHITE

    # Right Content Card
    tb_r = s1.shapes.add_textbox(Inches(5.6), Inches(1.4), Inches(6.8), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    
    pr_cat = tf_r.paragraphs[0]
    pr_cat.text = "PROYECTO DE DIGITALIZACIÓN APLICADA AL SISTEMA PRODUCTIVO"
    pr_cat.font.name = "Calibri"
    pr_cat.font.size = Pt(12)
    pr_cat.font.bold = True
    pr_cat.font.color.rgb = C_ACCENT
    pr_cat.space_after = Pt(12)
    
    pr_t = tf_r.add_paragraph()
    pr_t.text = "Plan Integral de Implementación de Tecnologías Habilitadoras Digitales (THD)"
    pr_t.font.name = "Calibri"
    pr_t.font.size = Pt(26)
    pr_t.font.bold = True
    pr_t.font.color.rgb = C_DARK
    pr_t.space_after = Pt(20)
    
    pr_desc = tf_r.add_paragraph()
    pr_desc.text = "Modernización de procesos internos, construcción de viviendas inteligentes con estándar Passivhaus y dinamización de zonas comunes (Smart Villages)."
    pr_desc.font.name = "Calibri"
    pr_desc.font.size = Pt(12)
    pr_desc.font.color.rgb = C_MUTED
    pr_desc.space_after = Pt(28)
    
    pr_meta = tf_r.add_paragraph()
    pr_meta.text = "Autor: Kevin  |  Curso: 2º DAM (Desarrollo de Aplicaciones Multiplataforma)\nConvocatoria Académica: 2025 - 2026  |  Entrega en Aules"
    pr_meta.font.name = "Calibri"
    pr_meta.font.size = Pt(12)
    pr_meta.font.bold = True
    pr_meta.font.color.rgb = C_MID

    # =========================================================================
    # DIAPOSITIVA 2: CONTEXTO Y OPORTUNIDAD ESTRATÉGICA
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Contexto Comarcal y Visión de Carlos y Esther")
    
    create_card(s2, Inches(0.9), Inches(1.8), Inches(3.6), Inches(5.0),
        "El Territorio y la Oportunidad",
        [
            "Comarca de Pequeños Pueblos: Separados por muy pocos kilómetros (3-8 km) en un entorno natural privilegiado.",
            "Llegada de Jóvenes: Profesionales que gracias al teletrabajo huyen de las grandes urbes buscando calidad de vida y calma.",
            "Exigencia de Servicios: Quieren naturaleza, pero sin renunciar a conectividad de alta velocidad y comodidades del siglo XXI.",
            "Compromiso de Origen: Carlos y Esther vuelven al valle donde se criaron para renovar el modelo constructivo."
        ],
        "Dinámica socioeconómica rural"
    )
    
    create_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "El Reto de la Vivienda Antigua",
        [
            "Parque Envejecido: La gran mayoría de viviendas son antiguas y presentan graves déficits térmicos y de salubridad.",
            "Necesidad de Rehabilitación: Muchas requieren una transformación integral para ser habitables hoy día.",
            "Apoyo Municipal: Ayuntamientos aprueban medidas para fomentar nueva edificación ecológica y respetuosa.",
            "Preservación Vernácula: Obligación de conservar la estética tradicional de piedra, madera y volumetrías históricas."
        ],
        "Condicionantes del entorno edificado"
    )

    create_card(s2, Inches(8.7), Inches(1.8), Inches(3.7), Inches(5.0),
        "La Solución: El Showroom Piloto",
        [
            "Modelo Sostenible y Económico: Casas industrializadas con estándar Passivhaus y consumo de energía casi nulo.",
            "Construcción Piloto (90 m²): Sede técnica de EcoRural y vivienda demostradora abierta a clientes y autoridades.",
            "Living Lab Activo: Espacio donde validar in situ la domótica, el confort térmico y la autogeneración solar.",
            "Enfoque Global: La digitalización no queda dentro de la empresa; se expande a la vivienda y al pueblo."
        ],
        "Estrategia de implantación gradual"
    )

    # =========================================================================
    # DIAPOSITIVA 3: TAREA 1 - THD EN PROCESOS INTERNOS
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Tarea 1: Digitalización de Procesos Internos en EcoRural")
    
    create_card(s3, Inches(0.9), Inches(1.8), Inches(5.6), Inches(2.4),
        "BIM (5D/6D) y Gemelo Digital",
        [
            "Modelado 3D inteligente: Integración de geometría, costes (5D) y sostenibilidad (6D).",
            "Detección de interferencias: Elimina errores entre tuberías y estructura antes de construir.",
            "Impacto: Ahorro de un 30% en sobrecostes y mermas en compras de materiales."
        ]
    )
    
    create_card(s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.4),
        "Cloud Computing y CDE Colaborativo",
        [
            "Gestión centralizada: Planos y modelos sincronizados en la nube en tiempo real.",
            "Movilidad en obra: Acceso inmediato desde tablets por parte de arquitectos y gremios.",
            "Impacto: Desaparición total del papel y cero desplazamientos innecesarios."
        ]
    )
    
    create_card(s3, Inches(0.9), Inches(4.4), Inches(5.6), Inches(2.4),
        "Drones Topográficos y Termográficos",
        [
            "Fotogrametría aérea: Levantamiento 3D de parcelas escarpadas en 2 horas en vez de días.",
            "Inspección de cubiertas: Detección de fisuras y fugas de calor sin subir a andamios.",
            "Impacto: Supresión del riesgo laboral de caídas y documentación gráfica para clientes."
        ]
    )
    
    create_card(s3, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.4),
        "IA Generativa, Ciberseguridad y ERP Cloud",
        [
            "Diseño bioclimático por IA: Orientación y sombras optimizadas según el microclima.",
            "Seguridad Zero Trust + Backup 3-2-1: Protección de planos frente a ransomware.",
            "ERP/CRM en Nube: Facturación electrónica automática (VeriFactu) y margen en tiempo real."
        ]
    )

    # =========================================================================
    # DIAPOSITIVA 4: TAREA 2 - THD EN LAS VIVIENDAS
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Tarea 2: THD Aplicables en la Construcción y Viviendas")
    
    create_card(s4, Inches(0.9), Inches(1.8), Inches(3.6), Inches(5.0),
        "IoT y Domótica Abierta",
        [
            "Protocolo Abierto Matter: Interoperabilidad sin depender de marcas cerradas.",
            "Salubridad Interior: Monitorización continua de CO2, humedad y compuestos orgánicos (COVs).",
            "Ventilación Mecánica (VMC): Recuperación del 90% del calor con aire siempre filtrado.",
            "Protección Antifugas: Electroválvulas de corte de agua automático ante roturas en tuberías.",
            "Beneficio: Confort térmico idóneo para teletrabajo y ahorro del 30% en energía."
        ],
        "Ambiente saludable y eficiente"
    )

    create_card(s4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Smart Energy e IA Predictiva",
        [
            "Solar Integrada: Tejas solares que respetan la estética tradicional de pizarra.",
            "Baterías Inteligentes (LiFePO4): Acumulación eléctrica de 10-15 kWh para la noche.",
            "Algoritmo Predictivo IA: Cruza tarifas horarias y previsión climática de lluvia o sol.",
            "Aerotermia de Alta Eficiencia: Suelo radiante y agua caliente con COP superior a 4.",
            "Beneficio: Facturas casi nulas y blindaje ante cortes de red eléctrica en el campo."
        ],
        "Vivienda de consumo casi nulo (nZEB)"
    )

    create_card(s4, Inches(8.7), Inches(1.8), Inches(3.7), Inches(5.0),
        "Vanguardia: 3D, VR y Blockchain",
        [
            "Impresión 3D: Fabricación de celosías tradicionales y remates de cubierta sin residuo.",
            "Showroom VR: Los clientes pasean y eligen los acabados interiores antes de construir.",
            "Realidad Aumentada (AR): Los operarios visualizan canalizaciones ocultas en la pared.",
            "Blockchain y Pasaporte Digital: Certificación inalterable de maderas locales y mantenimiento.",
            "Biometría: Apertura por huella y códigos temporales para paquetería mientras trabajas."
        ],
        "Transparencia y personalización total"
    )

    # =========================================================================
    # DIAPOSITIVA 5: TAREA 3 - THD EN LOS PUEBLOS (SMART VILLAGE)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Tarea 3: Modernización Comarcal y 'Smart Village'")
    
    create_card(s5, Inches(0.9), Inches(1.8), Inches(5.6), Inches(2.4),
        "Alumbrado Público Inteligente (IoT)",
        [
            "Luminarias LED con sensor de presencia: Modo reposo (20%) si no hay viandantes ni tráfico.",
            "Diseño rústico tradicional: Estética de forja que aloja en su interior la electrónica IoT.",
            "Impacto: Ahorro de más del 70% en luz municipal y eliminación de la contaminación lumínica (Cielos Starlight)."
        ]
    )

    create_card(s5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.4),
        "Gestión de Residuos y Riego Eficiente",
        [
            "Contenedores sensorizados: Ultrasonidos que avisan a la Mancomunidad al llegar al 80% de llenado.",
            "Riego predictivo en parques: Sondas de humedad en suelo que bloquean el agua si va a llover.",
            "Impacto: Reducción del 45% en consumo de gasóleo de camiones y del 50% de agua en zonas verdes."
        ]
    )

    create_card(s5, Inches(0.9), Inches(4.4), Inches(5.6), Inches(2.4),
        "Red Comarcal LoRaWAN y Wi-Fi Público",
        [
            "Cobertura de bajo consumo: 3 antenas que cubren todos los valles para telemetría municipal.",
            "Plazas conectadas (WiFi4EU): Permite a teletrabajadores y vecinos trabajar en plazas al aire libre.",
            "Impacto: Red abierta que moderniza también a agricultores y ganaderos locales (ganado con GPS)."
        ]
    )

    create_card(s5, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.4),
        "Alerta de Incendios y Plataforma Ciudadana",
        [
            "Detección térmica con IA: Alerta de conatos de fuego forestal en menos de 5 minutos.",
            "Boyas en ríos: Monitorización en tiempo real para prever riadas o crecidas repentinas.",
            "App Municipal + AR: Envío de averías con foto/GPS y rutas turísticas culturales con Realidad Aumentada."
        ]
    )

    # =========================================================================
    # DIAPOSITIVA 6: TAREA 4 - EVALUACIÓN DE IMPACTO AMBIENTAL
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. Tarea 4: Evaluación de Impacto Medioambiental y Circularidad")
    
    create_card(s6, Inches(0.9), Inches(1.8), Inches(3.6), Inches(5.0),
        "Descarbonización y Energía",
        [
            "Estándar Passivhaus: Demanda de calefacción inferior a 15 kWh/m² al año.",
            "Ahorro de Emisiones: Cada vivienda evita 4,2 toneladas de CO2 al año gracias a solar y aerotermia.",
            "Eficiencia Global: La gestión predictiva con IA evita consumir energía en las horas más contaminantes de la red.",
            "Movilidad Limpia: Integración de cargador de vehículo eléctrico alimentado con excedentes solares."
        ],
        "Hacia la huella de carbono negativa"
    )

    create_card(s6, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Economía Circular y KM Cero",
        [
            "Materiales Autóctonos: Madera de bosques comarcales sostenibles y mampostería de piedra local.",
            "Reutilización en Obra: En rehabilitación, se recupera el 80% de la piedra existente mediante escaneo BIM.",
            "Fabricación Modular: Las piezas preensambladas reducen los escombros a pie de obra en un 80%.",
            "Tratamiento de Aguas: Saneamiento con fitodepuración natural (humedales) sin químicos vertidos a ríos."
        ],
        "Construcción regenerativa y local"
    )

    create_card(s6, Inches(8.7), Inches(1.8), Inches(3.7), Inches(5.0),
        "Protección del Ecosistema",
        [
            "Cielos Nocturnos Limpios: La telegestión de farolas protege la fauna nocturna y permite el astroturismo.",
            "Ahorro de Agua Municipal: Las zonas comunes reducen a la mitad su consumo de riego con sondas de humedad.",
            "Resiliencia Forestal: La red de cámaras térmicas previene grandes desastres de incendios en el monte.",
            "Equilibrio Visual: Las casas se mimetizan estéticamente con el paisaje sin estridencias."
        ],
        "Sinergia con la naturaleza rural"
    )

    # =========================================================================
    # DIAPOSITIVA 7: TAREA 4 - IMPACTO SOCIAL Y MITIGACIÓN DE RIESGOS
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "6. Tarea 4: Impacto Social, Inclusión y Matriz de Mitigación")
    
    create_card(s7, Inches(0.9), Inches(1.8), Inches(5.6), Inches(5.0),
        "Impacto Social: Freno a la Despoblación",
        [
            "Fijación de Jóvenes: La oferta de casas tecnológicas asequibles atrae a nómadas digitales y jóvenes familias que reactivan escuelas y comercios.",
            "Capacitación de Gremios Locales: EcoRural forma a albañiles, fontaneros y electricistas del valle en domótica y tecnologías verdes.",
            "Silver Economy y Mayores: Domótica asistencial invisible (luz nocturna suave que evita caídas, teleasistencia por voz, sensores de inactividad) para que los ancianos vivan seguros en su hogar de siempre.",
            "Cohesión Comunitaria: La app comarcal y los espacios públicos conectados fortalecen la vida colectiva del valle."
        ],
        "Beneficios transformadores en la comarca"
    )

    create_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "Matriz de Riesgos y Mitigación",
        [
            "Riesgo de Ruptura Estética: Miedo a que la tecnología rompa el encanto tradicional.\n→ Mitigación: Tecnología oculta, tejas solares de pizarra y farolas rústicas de forja con electrónica embutida.",
            "Riesgo de Brecha Digital: Dificultad para vecinos mayores con la tecnología.\n→ Mitigación: Automatismos pasivos sin pantallas, pulsadores de toda la vida y talleres comunitarios de Carlos y Esther.",
            "Riesgo de Gentrificación y Subida de Precios: Expulsión de la juventud local.\n→ Mitigación: Precios contenidos por industrialización y convenios de suelo público con los consistorios.",
            "Riesgo de Basura Electrónica (e-waste): Dispositivos obsoletos en pocos años.\n→ Mitigación: Protocolos abiertos (Matter) y reciclaje en el Punto Limpio comarcal."
        ],
        "Gestión preventiva y responsable"
    )

    # =========================================================================
    # DIAPOSITIVA 8: PLAN DE IMPLANTACIÓN Y CONCLUSIONES
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "7. Hoja de Ruta, Conclusiones y Defensa")
    
    create_card(s8, Inches(0.9), Inches(1.8), Inches(3.6), Inches(5.0),
        "Fases de Implantación",
        [
            "Fase 1 (M1 - M4): Digitalización Interna\n• BIM 5D, Nube CDE, Dron y Ciberseguridad.\n• Inversión: 18.500 € (Kit Digital / Fondos UE).",
            "Fase 2 (M5 - M10): Vivienda Piloto\n• Construcción del Showroom y Living Lab.\n• Integración de IoT Matter, baterías y VR.\n• Inversión: 125.000 € (Crédito verde).",
            "Fase 3 (M11 - M18): Despliegue Smart Village\n• Venta de promociones y red LoRaWAN comarcal.\n• Convenios de telegestión con ayuntamientos."
        ],
        "Cronograma y presupuesto controlado"
    )

    create_card(s8, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "Valoración Estratégica",
        [
            "Sin THD no hay Sostenibilidad Real: Las THD permiten medir la energía, optimizar materiales y garantizar el impacto neto cero.",
            "Alta Rentabilidad Económica: La eliminación del 30% en errores de obra y los plazos acortados amortizan la inversión desde el 1er año.",
            "Modelo Exportable: EcoRural demuestra que la innovación es el escudo más potente para el futuro del medio rural.",
            "Conexión con el Territorio: Se preserva la memoria de los pueblos modernizando sus infraestructuras."
        ],
        "Lecciones del plan de Carlos y Esther"
    )

    create_card(s8, Inches(8.7), Inches(1.8), Inches(3.7), Inches(5.0),
        "Datos de Defensa y Aules",
        [
            "Proyecto: Plan de Transformación Digital de EcoRural Construcciones.",
            "Módulo: Digitalización Aplicada al Sistema Productivo.",
            "Alumno: Kevin (2º DAM).",
            "Documentación Entregada:\n• Informe Técnico Oficial (Calibri 11, justificado, 1.5, tablas completas).\n• Presentación Ejecutiva en Diapositivas.\n• Enlace y entrega formal en Aules.",
            "¡Muchas gracias por su atención! Turno de preguntas y defensa."
        ],
        "Cierre institucional y académico"
    )

    # Save presentation
    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(output_dir, "Presentacion_THD_EcoRural_Kevin.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")
    return output_path

if __name__ == "__main__":
    create_presentation()
