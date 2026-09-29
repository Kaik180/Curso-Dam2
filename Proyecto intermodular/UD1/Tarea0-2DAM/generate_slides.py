import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    curr_dir = os.path.dirname(__file__)

    # Memorix Color Palette
    C_NAVY = RGBColor(5, 12, 38)          # #050C26
    C_MINT = RGBColor(0, 180, 150)        # #00B496
    C_BLUE = RGBColor(0, 56, 255)         # #0038FF
    C_CORAL = RGBColor(224, 86, 36)       # #E05624
    C_BG_SLIDE = RGBColor(248, 250, 252)  # #F8FAFC
    C_CARD_BG = RGBColor(255, 255, 255)   # #FFFFFF
    C_BORDER = RGBColor(226, 232, 240)    # #E2E8F0
    C_TEXT_DARK = RGBColor(35, 47, 62)    # #232F3E
    C_TEXT_MUTED = RGBColor(100, 116, 139) # #64748B

    def draw_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG_SLIDE
        bg.line.color.rgb = C_BG_SLIDE
        return bg

    def add_top_nav(slide, tag_text, slide_num_str):
        # Top banner line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_MINT
        top_bar.line.color.rgb = C_MINT

        # Header tag
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(8), Inches(0.45))
        p = tb.text_frame.paragraphs[0]
        p.text = "Memorix  •  Proyecto Intermodular 2º DAM"
        p.font.name = "Calibri"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_MINT

        # Slide indicator on right
        tb_num = slide.shapes.add_textbox(Inches(9.5), Inches(0.2), Inches(3), Inches(0.45))
        p_num = tb_num.text_frame.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = slide_num_str
        p_num.font.name = "Calibri"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = C_TEXT_MUTED

    def add_slide_header(slide, title, subtitle):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = 0
        tf.margin_left = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Calibri"
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # DIAPOSITIVA 1: EL PROBLEMA Y A QUIÉN AFECTA
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    draw_bg(slide1)
    add_top_nav(slide1, "DIAGNÓSTICO DEL PROBLEMA", "01 / 03  •  EL PROBLEMA")
    add_slide_header(slide1, "1. Diagnóstico del Problema y Colectivos Afectados", "¿Qué ocurre en la vida cotidiana y a quién afecta directamente?")

    # Card 1: ¿Qué ocurre?
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_CARD_BG
    card1.line.color.rgb = C_BORDER
    card1.line.width = Pt(1.5)

    tb_c1 = slide1.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.1))
    tf1 = tb_c1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "🚨 ¿QUÉ OCURRE?"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    items_c1 = [
        ("Olvidos y Confusión en la Medicación", "Casi el 50% de pacientes crónicos no toma bien su tratamiento (OMS), provocando recaídas y visitas a urgencias."),
        ("Dificultad con Posologías y Prospectos", "La letra diminuta y las explicaciones técnicas confunden al paciente sobre cómo y cuándo tomar cada pastilla."),
        ("Abandono de la Rehabilitación en Casa", "Las hojas de papel con ejercicios para dolores musculares o traumatismos se extravían y no guían la postura correcta.")
    ]
    for tit, desc in items_c1:
        p_t = tf1.add_paragraph()
        p_t.text = f"• {tit}: "
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY
        p_t.space_after = Pt(1)

        p_d = tf1.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = C_TEXT_DARK
        p_d.space_after = Pt(10)

    # Card 2: ¿A quién afecta?
    card2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.5))
    card2.fill.solid()
    card2.fill.fore_color.rgb = C_CARD_BG
    card2.line.color.rgb = C_BORDER
    card2.line.width = Pt(1.5)

    tb_c2 = slide1.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.1))
    tf2 = tb_c2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "👥 ¿A QUIÉN AFECTA?"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_after = Pt(12)

    items_c2 = [
        ("Personas Mayores y Pacientes con Olvidos", "Especialmente vulnerables por polimedicación o principios de Alzheimer y deterioro cognitivo leve."),
        ("Familiares y Cuidadores", "Viven con la preocupación y el estrés diario de tener que estar comprobando si se han tomado la medicación."),
        ("Personas con Lesiones Musculares o Dolor", "Pacientes en recuperación de traumatismos, contracturas o con dolores de espalda que necesitan guías de ejercicios.")
    ]
    for tit, desc in items_c2:
        p_t = tf2.add_paragraph()
        p_t.text = f"• {tit}: "
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY
        p_t.space_after = Pt(1)

        p_d = tf2.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = C_TEXT_DARK
        p_d.space_after = Pt(10)

    # Footer banner slide 1
    b_foot1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.55))
    b_foot1.fill.solid()
    b_foot1.fill.fore_color.rgb = C_NAVY
    b_foot1.line.color.rgb = C_NAVY
    p_bf1 = b_foot1.text_frame.paragraphs[0]
    p_bf1.text = "💡 IMPACTO: Alineado con ODS 3 (Salud y Bienestar) y ODS 10 (Reducción de las Desigualdades)"
    p_bf1.font.name = "Calibri"
    p_bf1.font.size = Pt(11)
    p_bf1.font.bold = True
    p_bf1.font.color.rgb = RGBColor(255, 255, 255)
    p_bf1.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # DIAPOSITIVA 2: LA SOLUCIÓN MEMORIX Y FUNCIONALIDADES
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    draw_bg(slide2)
    add_top_nav(slide2, "PROPUESTA DIGITAL", "02 / 03  •  LA SOLUCIÓN")
    add_slide_header(slide2, "2. Memorix: Propuesta Digital y Funcionalidades", "Asistente móvil accesible para el control de medicamentos y ejercicios de recuperación")

    # Image Left
    img_mockup = os.path.join(curr_dir, "mockup_memorix.jpg")
    if os.path.exists(img_mockup):
        slide2.shapes.add_picture(img_mockup, Inches(0.8), Inches(1.8), width=Inches(5.5))

    # Card Right
    card_sol = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.8), Inches(5.9), Inches(4.5))
    card_sol.fill.solid()
    card_sol.fill.fore_color.rgb = C_CARD_BG
    card_sol.line.color.rgb = C_BORDER
    card_sol.line.width = Pt(1.5)

    tb_sol = slide2.shapes.add_textbox(Inches(6.9), Inches(1.95), Inches(5.3), Inches(4.2))
    tf_sol = tb_sol.text_frame
    tf_sol.word_wrap = True

    p = tf_sol.paragraphs[0]
    p.text = "⚡ FUNCIONALIDADES CLAVE"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p.space_after = Pt(10)

    funcs_list = [
        ("🔊 Alarmas con Locución por Voz", "Lee en voz alta el nombre de la pastilla, la dosis y cómo tomarla (ej. con comida)."),
        ("👆 Confirmación Fácil de Un Toque", "Botón táctil grande para registrar la toma y evitar tomar la pastilla dos veces por despiste."),
        ("🛡️ Aviso Automático al Cuidador", "Manda una alerta al móvil del familiar si pasan 20 minutos sin confirmarse la toma."),
        ("🏋️ Ejercicios Musculares Guiados", "Rutinas para dolores o traumatismos con temporizador, series y descansos en pantalla."),
        ("🎨 Interfaz de Alto Contraste", "Tipografía clara y colores contrastados pensados para personas mayores o con poca vista.")
    ]
    for f_tit, f_desc in funcs_list:
        p_ft = tf_sol.add_paragraph()
        p_ft.text = f_tit
        p_ft.font.name = "Calibri"
        p_ft.font.size = Pt(11.5)
        p_ft.font.bold = True
        p_ft.font.color.rgb = C_MINT
        p_ft.space_after = Pt(1)

        p_fd = tf_sol.add_paragraph()
        p_fd.text = f_desc
        p_fd.font.name = "Calibri"
        p_fd.font.size = Pt(10)
        p_fd.font.color.rgb = C_TEXT_DARK
        p_fd.space_after = Pt(7)

    # Footer banner slide 2
    b_foot2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.55))
    b_foot2.fill.solid()
    b_foot2.fill.fore_color.rgb = C_NAVY
    b_foot2.line.color.rgb = C_NAVY
    p_bf2 = b_foot2.text_frame.paragraphs[0]
    p_bf2.text = "📱 Solución Multiplataforma (Android / iOS) • Modo Fácil Diseñado para la Tercera Edad"
    p_bf2.font.name = "Calibri"
    p_bf2.font.size = Pt(11)
    p_bf2.font.bold = True
    p_bf2.font.color.rgb = RGBColor(255, 255, 255)
    p_bf2.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # DIAPOSITIVA 3: VALOR AÑADIDO Y MEJORA FUTURA
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    draw_bg(slide3)
    add_top_nav(slide3, "VALOR E INNOVACIÓN", "03 / 03  •  VALOR AÑADIDO")
    add_slide_header(slide3, "3. Valor Diferencial y Mejoras Futuras", "¿Por qué resulta útil esta solución y cómo puede evolucionar?")

    # Card Left: Valor Añadido
    card_va = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.5))
    card_va.fill.solid()
    card_va.fill.fore_color.rgb = C_CARD_BG
    card_va.line.color.rgb = C_BORDER
    card_va.line.width = Pt(1.5)

    tb_va = slide3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.1))
    tf_va = tb_va.text_frame
    tf_va.word_wrap = True

    p = tf_va.paragraphs[0]
    p.text = "🌟 VALOR AÑADIDO (¿POR QUÉ ES ÚTIL?)"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_MINT
    p.space_after = Pt(12)

    val_items = [
        ("Solución 2 en 1 (Medicamentos + Ejercicio)", "Une en una sola aplicación el control de los tratamientos y la recuperación física activa."),
        ("Mayor Autonomía para Personas Mayores", "Las explicaciones por voz y los botones grandes permiten manejarse sin depender siempre de otra persona."),
        ("Tranquilidad para los Familiares", "Supervisión sin agobios; solo reciben alerta en el móvil si su familiar no confirma una toma."),
        ("Constancia en la Rehabilitación", "Las guías visuales con cronómetro ayudan a no abandonar los ejercicios musculares en casa.")
    ]
    for v_tit, v_desc in val_items:
        p_t = tf_va.add_paragraph()
        p_t.text = f"✔ {v_tit}: "
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(11.5)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY
        p_t.space_after = Pt(1)

        p_d = tf_va.add_paragraph()
        p_d.text = f"   {v_desc}"
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = C_TEXT_DARK
        p_d.space_after = Pt(8)

    # Card Right: Mejoras Futuras
    card_mf = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.5))
    card_mf.fill.solid()
    card_mf.fill.fore_color.rgb = C_CARD_BG
    card_mf.line.color.rgb = C_BORDER
    card_mf.line.width = Pt(1.5)

    tb_mf = slide3.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.1))
    tf_mf = tb_mf.text_frame
    tf_mf.word_wrap = True

    p = tf_mf.paragraphs[0]
    p.text = "🚀 MEJORAS FUTURAS PREVISTAS"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_after = Pt(12)

    fut_items = [
        ("Visión Artificial con Cámara Móvil", "Detección de la postura del usuario en tiempo real durante los ejercicios para avisar por voz si la espalda está mal colocada."),
        ("Conexión con Pulseras Inteligentes (Smartwatch)", "Avisos de medicación por vibración en la muñeca y detección de caídas accidentales en ancianos."),
        ("Escaneo de Cajas de Medicamentos", "Reconocimiento automático de la pastilla con la cámara para añadir la alarma sin tener que escribir.")
    ]
    for f_tit, f_desc in fut_items:
        p_t = tf_mf.add_paragraph()
        p_t.text = f"✦ {f_tit}: "
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(11.5)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY
        p_t.space_after = Pt(1)

        p_d = tf_mf.add_paragraph()
        p_d.text = f"   {f_desc}"
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = C_TEXT_DARK
        p_d.space_after = Pt(8)

    # Footer banner slide 3
    b_foot3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.55))
    b_foot3.fill.solid()
    b_foot3.fill.fore_color.rgb = C_NAVY
    b_foot3.line.color.rgb = C_NAVY
    p_bf3 = b_foot3.text_frame.paragraphs[0]
    p_bf3.text = "🎯 MEMORIX  •  Proyecto Intermodular 2º DAM  •  Kevin Muñoz Córcoles"
    p_bf3.font.name = "Calibri"
    p_bf3.font.size = Pt(11)
    p_bf3.font.bold = True
    p_bf3.font.color.rgb = RGBColor(255, 255, 255)
    p_bf3.alignment = PP_ALIGN.CENTER

    output_pptx = os.path.join(curr_dir, "Tarea0_Memorix_Presentacion.pptx")
    prs.save(output_pptx)
    print(f"Presentation saved successfully at: {output_pptx}")

if __name__ == "__main__":
    create_presentation()
