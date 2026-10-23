import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

curr_dir = r"c:\Users\kevin\OneDrive\Documentos\GitHub\Curso-Dam2\Proyecto intermodular\UD1\Tarea0-2DAM"

def create_word_tarea0():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.different_first_page_header_footer = True

        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Proyecto Intermodular - 2º DAM  |  Tarea 0: LeadCraft AI")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(9)
        r_hdr.font.color.rgb = RGBColor(120, 144, 156)

        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr_l = p_ftr.add_run("LeadCraft AI  •  Kevin Muñoz Córcoles\t\tPágina ")
        r_ftr_l.font.name = "Calibri"
        r_ftr_l.font.size = Pt(9)
        r_ftr_l.font.color.rgb = RGBColor(120, 144, 156)
        
        fld_run = p_ftr.add_run()
        fld_run.font.name = "Calibri"
        fld_run.font.size = Pt(9)
        fld_run.font.color.rgb = RGBColor(120, 144, 156)
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        fld_run._r.append(fldChar1)
        fld_run._r.append(instrText)
        fld_run._r.append(fldChar2)
        fld_run._r.append(fldChar3)

    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(12)
    font_normal.color.rgb = RGBColor(35, 47, 62)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(3)

    bookmark_id = 0
    def add_bookmark(paragraph, name):
        nonlocal bookmark_id
        bookmark_id += 1
        b_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{name}"/>')
        b_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>')
        paragraph._p.insert(0, b_start)
        paragraph._p.append(b_end)

    def add_external_link(paragraph, url, text):
        part = paragraph.part
        r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} {nsdecls("r")} r:id="{r_id}" w:history="1"/>')
        new_run = parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr>'
            f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
            f'<w:color w:val="0B7285"/>'
            f'<w:u w:val="single"/>'
            f'</w:rPr><w:t>{text}</w:t></w:r>'
        )
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)

    def add_toc_item(paragraph, title, anchor_name, page_num_str):
        hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor_name}" w:history="1"/>')
        run_title = parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr>'
            f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
            f'<w:color w:val="0A192F"/>'
            f'<w:u w:val="single"/>'
            f'</w:rPr><w:t>{title}</w:t></w:r>'
        )
        hyperlink.append(run_title)
        paragraph._p.append(hyperlink)

        r_dots = paragraph.add_run(f"  {'.' * 44}  ")
        r_dots.font.color.rgb = RGBColor(160, 174, 192)
        r_dots.font.name = "Calibri"
        r_dots.font.size = Pt(11)

        r_p = paragraph.add_run(page_num_str)
        r_p.font.name = "Calibri"
        r_p.font.size = Pt(11)
        r_p.font.bold = True
        r_p.font.color.rgb = RGBColor(10, 25, 47)

    def add_section_heading(title_text, anchor_name=None, space_before=6, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(title_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(10, 25, 47)
        
        if anchor_name:
            add_bookmark(p, anchor_name)
        return p

    def add_body_p(text, bold_prefix=None, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(space_after)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = 'Calibri'
            r_b.font.size = Pt(12)
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(0, 150, 136)
        r_t = p.add_run(text)
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(12)
        r_t.font.color.rgb = RGBColor(40, 50, 60)
        return p

    # ========================== PÁGINA 1: PORTADA ==========================
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(40)

    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_after = Pt(8)
    r_b = p_badge.add_run("PROYECTO INTERMODULAR • 2º DAM")
    r_b.font.name = "Calibri"
    r_b.font.size = Pt(11)
    r_b.font.bold = True
    r_b.font.color.rgb = RGBColor(0, 180, 150)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.line_spacing = 1.1
    r_tit = p_title.add_run("LeadCraft AI")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(32)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(10, 25, 47)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    p_sub.paragraph_format.line_spacing = 1.2
    r_sub = p_sub.add_run("Servicio Integral de Gestión y Presencia Online para Ayudar al Pequeño Comercio Local a Llevar su Negocio en Internet")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0, 56, 255)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(16)
    r_d = p_div.add_run("—" * 30)
    r_d.font.color.rgb = RGBColor(203, 213, 225)

    p_desc_cov = doc.add_paragraph()
    p_desc_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc_cov.paragraph_format.space_after = Pt(100)
    p_desc_cov.paragraph_format.line_spacing = 1.3
    r_dc = p_desc_cov.add_run("Tarea 0: Una Solución Digital para un Problema Real\nPropuesta Inicial y Definición del Proyecto Intermodular")
    r_dc.font.name = "Calibri"
    r_dc.font.size = Pt(11)
    r_dc.font.color.rgb = RGBColor(100, 116, 139)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.line_spacing = 1.3
    p_meta.paragraph_format.space_after = Pt(0)
    
    r_aut_lbl = p_meta.add_run("Nombre y apellidos: ")
    r_aut_lbl.font.bold = True
    r_aut_lbl.font.size = Pt(12)
    r_aut_val = p_meta.add_run("Kevin Muñoz Córcoles\n")
    r_aut_val.font.size = Pt(12)
    
    r_cur_lbl = p_meta.add_run("Curso: ")
    r_cur_lbl.font.bold = True
    r_cur_lbl.font.size = Pt(12)
    r_cur_val = p_meta.add_run("2º Desarrollo de Aplicaciones Multiplataforma (DAM)\n")
    r_cur_val.font.size = Pt(12)

    r_mod_lbl = p_meta.add_run("Módulo: ")
    r_mod_lbl.font.bold = True
    r_mod_lbl.font.size = Pt(12)
    r_mod_val = p_meta.add_run("Proyecto Intermodular | Convocatoria 2026/2027")
    r_mod_val.font.size = Pt(12)

    doc.add_page_break()

    # ========================== PÁGINA 2: CONTENIDO 1 ==========================
    add_section_heading("Índice de Contenidos", "toc", space_before=0, space_after=3)
    
    toc_items = [
        ("1. Problema o Necesidad Detectada", "sec_problema", "Pág. 2"),
        ("2. Propuesta de Solución Digital: El Servicio LeadCraft AI", "sec_solucion", "Pág. 2"),
        ("3. Funcionalidades del Sistema", "sec_funcionalidades", "Pág. 3"),
        ("4. Mejora Futura y Escalabilidad", "sec_mejora", "Pág. 3"),
        ("5. Prototipo e Identidad Visual", "sec_imagen", "Pág. 3"),
        ("6. Fuentes y Referencias Consultadas", "sec_fuentes", "Pág. 3"),
    ]
    for title, anchor, pag in toc_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.line_spacing = 1.15
        p_item.paragraph_format.space_after = Pt(2)
        add_toc_item(p_item, title, anchor, pag)

    add_section_heading("1. Problema o Necesidad Detectada", "sec_problema", space_before=8, space_after=2)
    
    add_body_p(
        "En nuestras ciudades y barrios existen miles de pequeños comercios (peluquerías, talleres mecánicos, clínicas dentales o de fisioterapia y tiendas locales) que se ven completamente sobrepasados para llevar su presencia en Internet. La mayoría tienen una ficha en Google Maps, pero carecen de página web propia, no disponen de un sistema para reservar citas online y dependen de atender llamadas telefónicas mientras despachan a clientes en el mostrador. Esto les hace perder ventas a diario frente a grandes franquicias. Contratar agencias tradicionales resulta inasumible (cobran más de 1.000€ iniciales más costosas cuotas mensuales), y los dueños no disponen de tiempo ni conocimientos técnicos para gestionarlo por su cuenta. Al mismo tiempo, los profesionales y jóvenes desarrolladores que ofrecen servicios digitales invierten semanas contactando 'a puerta fría' sin una propuesta visual construida que convenza al comerciante.",
        bold_prefix="• ¿Qué ocurre? "
    )

    add_body_p(
        "Afecta principalmente a los dueños de pequeños comercios y autónomos locales, que no pueden atender el mostrador físico y al mismo tiempo gestionar citas, mensajes y catálogos en la red. También perjudica a los vecinos y clientes locales, que no pueden reservar hora ni informarse fuera del horario comercial. Por último, afecta a jóvenes titulados de DAM y consultores freelance, que necesitan un sistema ágil para empaquetar y vender este servicio de gestión digital a los comercios de su entorno.",
        bold_prefix="• ¿A quién afecta? "
    )

    add_body_p(
        "Resolver esta brecha es indispensable para la viabilidad del comercio tradicional. Venderles un servicio llave en mano y asequible que les gestione su web y reservas online les permite competir en igualdad de condiciones con el comercio electrónico. Este proyecto impacta directamente en el ODS 8 (Trabajo decente y crecimiento económico), al fomentar la productividad de microempresas, y en el ODS 9 (Industria, innovación e infraestructura), al democratizar el acceso a la tecnología y la inteligencia artificial para negocios de proximidad.",
        bold_prefix="• ¿Por qué debería mejorarse? "
    )

    add_section_heading("2. Propuesta de Solución Digital: El Servicio LeadCraft AI", "sec_solucion", space_before=8, space_after=2)

    add_body_p(
        "LeadCraft AI es una plataforma de software diseñada para comercializar un servicio integral de digitalización y gestión online para comercios locales. El núcleo del proyecto no consiste únicamente en crear páginas web estáticas, sino en vender un servicio activo de la aplicación que ayuda al comerciante a llevar su negocio o tienda en Internet de forma continuada: incluye el despliegue de su web profesional generada mediante IA, un gestor interactivo de citas y reservas en tiempo real, catálogo de servicios actualizado, canal directo de pedidos por WhatsApp y un panel sencillo desde el que el comerciante consulta sus reservas diarias sin complicaciones técnicas ni costes desproporcionados.",
        bold_prefix="• Explicación de la idea: "
    )

    add_body_p(
        "La solución permite: 1) Explorar negocios en Google Maps filtrando por sector y zona geográfica; 2) Auditar en 3 segundos sus carencias digitales (si carecen de web, citas online o redes); 3) Generar con una API de IA una web completa y personalizada con la información de Maps (fotos, horarios, carta de servicios); 4) Activar un módulo interactivo de reservas y contacto por WhatsApp para que la tienda atienda clientes 24/7; 5) Facilitar al dueño un panel de control para gestionar citas diarias y precios; y 6) Empaquetar y vender esta solución bajo una suscripción mensual asequible de mantenimiento y soporte.",
        bold_prefix="• ¿Qué permitiría hacer? "
    )

    add_body_p(
        "El servicio está dirigido a los comerciantes y autónomos locales (los clientes del servicio, que obtienen una solución lista para gestionar su negocio online sin esfuerzo técnico), sus clientes finales (que reservan citas y consultan productos desde el móvil), y el operador o consultor comercial (estudiante de DAM o desarrollador) que prospecta la zona y administra la implantación del servicio.",
        bold_prefix="• ¿Quién utilizaría la solución? "
    )

    doc.add_page_break()

    # ========================== PÁGINA 3: CONTENIDO 2 ==========================
    add_section_heading("3. Funcionalidades del Sistema", "sec_funcionalidades", space_before=0, space_after=2)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Funcionalidad"
    hdr_cells[1].text = "Descripción"
    hdr_cells[0].width = Inches(2.2)
    hdr_cells[1].width = Inches(4.5)

    for cell in hdr_cells:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A192F"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(2.5)
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.bold = True
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(255, 255, 255)

    funcs_data = [
        ("Búsqueda Geográfica y Auditoría", "Localiza comercios sobre Google Maps y analiza en tiempo real sus carencias digitales (ausencia de web, citas online o datos desactualizados)."),
        ("Generador Web con IA para Negocios", "A partir de la ficha de Maps, una API de IA redacta textos comerciales, estructura los servicios y maqueta una web moderna y responsive."),
        ("Gestor Integrado de Citas y Reservas", "Incorpora un calendario interactivo con franjas horarias y confirmación inmediata para que los clientes reserven online 24/7."),
        ("Canal Directo de Contacto por WhatsApp", "Enlace automatizado que envía las consultas de clientes y solicitudes de servicio directamente al móvil del comerciante."),
        ("Panel de Control para el Comerciante", "Interfaz sencilla para consultar las reservas diarias, confirmar citas y actualizar precios u horarios desde cualquier dispositivo."),
        ("Dossier Comercial y Despliegue Cloud", "Genera informes descargables en PDF para la venta del servicio y permite exportar y publicar la web en servidores en la nube.")
    ]

    for idx, (func, desc) in enumerate(funcs_data):
        row_cells = table.add_row().cells
        row_cells[0].width = Inches(2.2)
        row_cells[1].width = Inches(4.5)
        row_cells[0].text = func
        row_cells[1].text = desc

        fill_color = "F1F5F9" if idx % 2 == 0 else "FFFFFF"
        for i_col, cell in enumerate(row_cells):
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
            cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(30, 41, 59)
                    if i_col == 0:
                        r.font.bold = True

    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="0A192F"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    add_section_heading("4. Mejora Futura y Escalabilidad", "sec_mejora", space_before=6, space_after=2)
    add_body_p(
        "Como ampliación futura se prevé integrar un asistente virtual de WhatsApp impulsado por IA para atender preguntas frecuentes de los clientes y gestionar citas de forma automática las 24 horas del día. Además, se incorporará la pasarela de pagos con Stripe para permitir el cobro de señales de reserva o pedidos online y la sincronización bidireccional automática con la agenda de Google Calendar del negocio.",
        bold_prefix="• Evolución prevista: "
    )

    add_section_heading("5. Prototipo e Identidad Visual de la Propuesta", "sec_imagen", space_before=6, space_after=2)
    
    img_mockup = os.path.join(curr_dir, "mockup_leadcraft.jpg")
    if os.path.exists(img_mockup):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(1)
        r_img = p_img.add_run()
        r_img.add_picture(img_mockup, width=Inches(3.8))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(4)
        r_cap = p_cap.add_run("Figura 1: Prototipo de LeadCraft AI mostrando el buscador geográfico de comercios locales y el panel de auditoría para la venta del servicio digital.")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    add_section_heading("6. Fuentes y Referencias Consultadas", "sec_fuentes", space_before=5, space_after=2)
    
    p_f1 = doc.add_paragraph()
    p_f1.paragraph_format.line_spacing = 1.15
    p_f1.paragraph_format.space_after = Pt(1.5)
    r1 = p_f1.add_run("• Instituto Nacional de Estadística (INE): ")
    r1.font.name = "Calibri"
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    add_external_link(p_f1, "https://www.ine.es", "Encuesta sobre el uso de TIC y comercio electrónico en las empresas españolas (Enlace web)")

    p_f2 = doc.add_paragraph()
    p_f2.paragraph_format.line_spacing = 1.15
    p_f2.paragraph_format.space_after = Pt(1.5)
    r2 = p_f2.add_run("• Google Cloud Platform: ")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10.5)
    r2.font.bold = True
    add_external_link(p_f2, "https://developers.google.com/maps/documentation/places/web-service/overview", "Documentación oficial de Google Places API (Enlace web)")

    p_f3 = doc.add_paragraph()
    p_f3.paragraph_format.line_spacing = 1.15
    p_f3.paragraph_format.space_after = Pt(1.5)
    r3 = p_f3.add_run("• Naciones Unidas: ")
    r3.font.name = "Calibri"
    r3.font.size = Pt(10.5)
    r3.font.bold = True
    add_external_link(p_f3, "https://www.un.org/sustainabledevelopment/es/economic-growth/", "Objetivo de Desarrollo Sostenible 8: Trabajo Decente y Crecimiento Económico (Enlace web)")

    out_docx = os.path.join(curr_dir, "Tarea0_LeadCraft_Kevin.docx")
    doc.save(out_docx)
    print("Word Tarea 0 LeadCraft saved at:", out_docx)

def create_slides_tarea0():
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)
    blank_layout = prs.slide_layouts[6]

    C_NAVY = PptxRGBColor(10, 25, 47)
    C_CYAN = PptxRGBColor(0, 210, 180)
    C_BLUE = PptxRGBColor(13, 110, 253)
    C_CORAL = PptxRGBColor(224, 86, 36)
    C_BG = PptxRGBColor(248, 250, 252)
    C_CARD_BG = PptxRGBColor(255, 255, 255)
    C_BORDER = PptxRGBColor(226, 232, 240)
    C_TEXT = PptxRGBColor(35, 47, 62)
    C_MUTED = PptxRGBColor(100, 116, 139)

    def draw_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.color.rgb = C_BG
        return bg

    def add_top_nav(slide, num_str):
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, PptxInches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_CYAN
        top_bar.line.color.rgb = C_CYAN

        tb = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.2), PptxInches(8), PptxInches(0.45))
        p = tb.text_frame.paragraphs[0]
        p.text = "LeadCraft AI  •  Proyecto Intermodular 2º DAM"
        p.font.name = "Calibri"
        p.font.size = PptxPt(12)
        p.font.bold = True
        p.font.color.rgb = C_CYAN

        tb_num = slide.shapes.add_textbox(PptxInches(9.5), PptxInches(0.2), PptxInches(3), PptxInches(0.45))
        p_num = tb_num.text_frame.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = num_str
        p_num.font.name = "Calibri"
        p_num.font.size = PptxPt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = C_MUTED

    def add_header(slide, title, subtitle):
        tb = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.65), PptxInches(11.7), PptxInches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = 0
        tf.margin_left = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Calibri"
        p1.font.size = PptxPt(24)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p1.space_after = PptxPt(2)

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Calibri"
        p2.font.size = PptxPt(13)
        p2.font.color.rgb = C_MUTED

    # SLIDE 1: PROBLEMA
    s1 = prs.slides.add_slide(blank_layout)
    draw_bg(s1)
    add_top_nav(s1, "01 / 03  •  EL PROBLEMA")
    add_header(s1, "1. Diagnóstico del Problema y Colectivos Afectados", "¿Por qué los comercios locales tienen dificultades para llevar su negocio en Internet?")

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(0.8), PptxInches(1.8), PptxInches(5.6), PptxInches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_CARD_BG
    card1.line.color.rgb = C_BORDER
    card1.line.width = PptxPt(1.5)

    tf1 = card1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🚨 ¿QUÉ OCURRE?"
    p.font.name = "Calibri"
    p.font.size = PptxPt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = PptxPt(12)

    items1 = [
        ("Sobrecarga en la Tienda Física", "Los comerciantes atienden a clientes en local y carecen de tiempo para llevar citas, web y catálogo online."),
        ("Costes Inasumibles de Agencias", "Las agencias tradicionales cobran más de 1.000€ más cuotas altas, algo inviable para un negocio de barrio."),
        ("Pérdida Continua de Clientes", "Sin reservas online ni presencia digital cuidada, pierden ventas a diario frente a grandes franquicias.")
    ]
    for tit, desc in items1:
        pt = tf1.add_paragraph()
        pt.text = f"• {tit}:"
        pt.font.name = "Calibri"
        pt.font.size = PptxPt(12)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pd = tf1.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.name = "Calibri"
        pd.font.size = PptxPt(11)
        pd.font.color.rgb = C_TEXT
        pd.space_after = PptxPt(10)

    card2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(6.8), PptxInches(1.8), PptxInches(5.7), PptxInches(4.5))
    card2.fill.solid()
    card2.fill.fore_color.rgb = C_CARD_BG
    card2.line.color.rgb = C_BORDER
    card2.line.width = PptxPt(1.5)

    tf2 = card2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "👥 ¿A QUIÉN AFECTA?"
    p.font.name = "Calibri"
    p.font.size = PptxPt(16)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_after = PptxPt(12)

    items2 = [
        ("Dueños de Tiendas y Comercios Locales", "Peluquerías, talleres, clínicas y bares que necesitan digitalizarse sin complicaciones técnicas."),
        ("Clientes y Vecinos del Barrio", "Buscan pedir cita, consultar precios y contactar desde el móvil 24/7 sin llamadas en horario comercial."),
        ("Titulados de DAM y Freelancers", "Oportunidad de ofrecer y vender un servicio recurrente de gestión digital 'llave en mano' en su entorno.")
    ]
    for tit, desc in items2:
        pt = tf2.add_paragraph()
        pt.text = f"• {tit}:"
        pt.font.name = "Calibri"
        pt.font.size = PptxPt(12)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pd = tf2.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.name = "Calibri"
        pd.font.size = PptxPt(11)
        pd.font.color.rgb = C_TEXT
        pd.space_after = PptxPt(10)

    b_foot1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptxInches(0.8), PptxInches(6.5), PptxInches(11.7), PptxInches(0.55))
    b_foot1.fill.solid()
    b_foot1.fill.fore_color.rgb = C_NAVY
    b_foot1.line.color.rgb = C_NAVY
    p_bf1 = b_foot1.text_frame.paragraphs[0]
    p_bf1.text = "💡 IMPACTO: Alineado con ODS 8 (Trabajo Decente y Crecimiento Económico) y ODS 9 (Innovación en Pymes)"
    p_bf1.font.name = "Calibri"
    p_bf1.font.size = PptxPt(11)
    p_bf1.font.bold = True
    p_bf1.font.color.rgb = PptxRGBColor(255, 255, 255)
    p_bf1.alignment = PP_ALIGN.CENTER

    # SLIDE 2: SOLUCIÓN
    s2 = prs.slides.add_slide(blank_layout)
    draw_bg(s2)
    add_top_nav(s2, "02 / 03  •  LA SOLUCIÓN")
    add_header(s2, "2. LeadCraft AI: Venta del Servicio Digital para Tiendas", "Servicio 'todo en uno' para poner en marcha y llevar el negocio del comerciante en Internet")

    img_path = os.path.join(curr_dir, "mockup_leadcraft.jpg")
    if os.path.exists(img_path):
        s2.shapes.add_picture(img_path, PptxInches(0.8), PptxInches(1.8), width=PptxInches(5.5))

    card_sol = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(6.6), PptxInches(1.8), PptxInches(5.9), PptxInches(4.5))
    card_sol.fill.solid()
    card_sol.fill.fore_color.rgb = C_CARD_BG
    card_sol.line.color.rgb = C_BORDER
    card_sol.line.width = PptxPt(1.5)

    tf_sol = card_sol.text_frame
    tf_sol.word_wrap = True
    p = tf_sol.paragraphs[0]
    p.text = "⚡ FUNCIONALIDADES CLAVE DEL SERVICIO"
    p.font.name = "Calibri"
    p.font.size = PptxPt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p.space_after = PptxPt(10)

    funcs_list = [
        ("🗺️ Prospección y Auditoría en Google Maps", "Localiza comercios y detecta carencias digitales (sin web, sin reservas)."),
        ("🤖 Web Asistida por IA para el Negocio", "Genera la web con horarios, catálogo y fotos extraídas de la ficha."),
        ("📅 Sistema Integrado de Citas Online 24/7", "Calendario interactivo con franjas horarias y confirmación inmediata."),
        ("💬 Canal de Contacto y Pedidos por WhatsApp", "Enlace directo para enviar consultas y reservas al móvil del comercio."),
        ("📱 Panel de Control Sencillo para el Dueño", "Gestión fácil de citas del día y cambios de horarios sin saber programar.")
    ]
    for ft, fd in funcs_list:
        pft = tf_sol.add_paragraph()
        pft.text = ft
        pft.font.name = "Calibri"
        pft.font.size = PptxPt(11.5)
        pft.font.bold = True
        pft.font.color.rgb = C_CYAN
        pfd = tf_sol.add_paragraph()
        pfd.text = fd
        pfd.font.name = "Calibri"
        pfd.font.size = PptxPt(10)
        pfd.font.color.rgb = C_TEXT
        pfd.space_after = PptxPt(6)

    b_foot2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptxInches(0.8), PptxInches(6.5), PptxInches(11.7), PptxInches(0.55))
    b_foot2.fill.solid()
    b_foot2.fill.fore_color.rgb = C_NAVY
    b_foot2.line.color.rgb = C_NAVY
    p_bf2 = b_foot2.text_frame.paragraphs[0]
    p_bf2.text = "🚀 Modelo B2B Micro-SaaS • Servicio Integral de Digitalización y Acompañamiento a Comercios"
    p_bf2.font.name = "Calibri"
    p_bf2.font.size = PptxPt(11)
    p_bf2.font.bold = True
    p_bf2.font.color.rgb = PptxRGBColor(255, 255, 255)
    p_bf2.alignment = PP_ALIGN.CENTER

    # SLIDE 3: VALOR AÑADIDO
    s3 = prs.slides.add_slide(blank_layout)
    draw_bg(s3)
    add_top_nav(s3, "03 / 03  •  VALOR AÑADIDO")
    add_header(s3, "3. Valor Añadido del Servicio y Mejoras Futuras", "¿Por qué resulta útil vender este servicio y cuál es su proyección futura?")

    card_va = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(0.8), PptxInches(1.8), PptxInches(5.6), PptxInches(4.5))
    card_va.fill.solid()
    card_va.fill.fore_color.rgb = C_CARD_BG
    card_va.line.color.rgb = C_BORDER
    card_va.line.width = PptxPt(1.5)

    tf_va = card_va.text_frame
    tf_va.word_wrap = True
    p = tf_va.paragraphs[0]
    p.text = "🌟 VALOR DEL SERVICIO (¿POR QUÉ ES ÚTIL?)"
    p.font.name = "Calibri"
    p.font.size = PptxPt(15)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.space_after = PptxPt(12)

    vals = [
        ("Servicio 'Llave en Mano'", "El comerciante no compra un software suelto, sino un servicio activo que le lleva su tienda en Internet."),
        ("Demostración Visual Instantánea", "El comercial se presenta con la tienda ya maquetada sobre datos de Google Maps, multiplicando el cierre de ventas."),
        ("Cuota Asequible por Suscripción", "Sustituye los presupuestos desorbitados de agencias por un servicio mensual accesible y mantenido."),
        ("Independencia para el Comerciante", "Panel intuitivo para consultar reservas sin depender de soporte técnico externo.")
    ]
    for vt, vd in vals:
        pt = tf_va.add_paragraph()
        pt.text = f"✔ {vt}:"
        pt.font.name = "Calibri"
        pt.font.size = PptxPt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pd = tf_va.add_paragraph()
        pd.text = f"   {vd}"
        pd.font.name = "Calibri"
        pd.font.size = PptxPt(10)
        pd.font.color.rgb = C_TEXT
        pd.space_after = PptxPt(8)

    card_mf = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(6.8), PptxInches(1.8), PptxInches(5.7), PptxInches(4.5))
    card_mf.fill.solid()
    card_mf.fill.fore_color.rgb = C_CARD_BG
    card_mf.line.color.rgb = C_BORDER
    card_mf.line.width = PptxPt(1.5)

    tf_mf = card_mf.text_frame
    tf_mf.word_wrap = True
    p = tf_mf.paragraphs[0]
    p.text = "🚀 MEJORAS FUTURAS PREVISTAS"
    p.font.name = "Calibri"
    p.font.size = PptxPt(15)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_after = PptxPt(12)

    futs = [
        ("Bot de WhatsApp con IA 24/7", "Atención automatizada para dudas frecuentes, confirmación de citas y recordatorios por chat."),
        ("Pagos Integrados con Stripe", "Cobro de señales o venta directa de productos y servicios desde la propia web."),
        ("Sincronización con Google Calendar", "Conexión bidireccional automática con la agenda personal y de empleados del comercio.")
    ]
    for ft, fd in futs:
        pt = tf_mf.add_paragraph()
        pt.text = f"✦ {ft}:"
        pt.font.name = "Calibri"
        pt.font.size = PptxPt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pd = tf_mf.add_paragraph()
        pd.text = f"   {fd}"
        pd.font.name = "Calibri"
        pd.font.size = PptxPt(10)
        pd.font.color.rgb = C_TEXT
        pd.space_after = PptxPt(8)

    b_foot3 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptxInches(0.8), PptxInches(6.5), PptxInches(11.7), PptxInches(0.55))
    b_foot3.fill.solid()
    b_foot3.fill.fore_color.rgb = C_NAVY
    b_foot3.line.color.rgb = C_NAVY
    p_bf3 = b_foot3.text_frame.paragraphs[0]
    p_bf3.text = "🎯 LEADCRAFT AI  •  Proyecto Intermodular 2º DAM  •  Kevin Muñoz Córcoles"
    p_bf3.font.name = "Calibri"
    p_bf3.font.size = PptxPt(11)
    p_bf3.font.bold = True
    p_bf3.font.color.rgb = PptxRGBColor(255, 255, 255)
    p_bf3.alignment = PP_ALIGN.CENTER

    out_pptx = os.path.join(curr_dir, "Tarea0_LeadCraft_Presentacion.pptx")
    prs.save(out_pptx)
    print("PowerPoint Tarea 0 LeadCraft saved at:", out_pptx)

if __name__ == "__main__":
    create_word_tarea0()
    create_slides_tarea0()
