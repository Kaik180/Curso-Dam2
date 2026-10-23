import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

curr_dir = r"c:\Users\kevin\OneDrive\Documentos\GitHub\Curso-Dam2\Proyecto intermodular\Hito1"

def build_word_hito1():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.different_first_page_header_footer = True

        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Proyecto Intermodular - 2º DAM  |  Hito 1: LeadCraft AI")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(120, 144, 156)

        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr_l = p_ftr.add_run("LeadCraft AI  •  Kevin Muñoz Córcoles\t\tPágina ")
        r_ftr_l.font.name = "Calibri"
        r_ftr_l.font.size = Pt(8.5)
        r_ftr_l.font.color.rgb = RGBColor(120, 144, 156)
        
        fld_run = p_ftr.add_run()
        fld_run.font.name = "Calibri"
        fld_run.font.size = Pt(8.5)
        fld_run.font.color.rgb = RGBColor(120, 144, 156)
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        fld_run._r.append(fldChar1)
        fld_run._r.append(instrText)
        fld_run._r.append(fldChar2)
        fld_run._r.append(fldChar3)

    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(10.5)
    font_normal.color.rgb = RGBColor(35, 47, 62)
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(2.5)

    bookmark_id = 0
    def add_bookmark(paragraph, name):
        nonlocal bookmark_id
        bookmark_id += 1
        b_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{name}"/>')
        b_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>')
        paragraph._p.insert(0, b_start)
        paragraph._p.append(b_end)

    def add_h1(text, anchor=None, space_before=5, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(10, 25, 47)
        if anchor:
            add_bookmark(p, anchor)
        return p

    def add_p(text, bold_prefix=None, space_after=2.5):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(space_after)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = 'Calibri'
            rb.font.size = Pt(10.5)
            rb.font.bold = True
            rb.font.color.rgb = RGBColor(10, 37, 64)
        rt = p.add_run(text)
        rt.font.name = 'Calibri'
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = RGBColor(40, 50, 60)
        return p

    # ========================== PÁGINA 1: PORTADA & ÍNDICE ==========================
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(20)

    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_after = Pt(4)
    rb = p_badge.add_run("PROYECTO INTERMODULAR • 2º DAM")
    rb.font.name = "Calibri"
    rb.font.size = Pt(10.5)
    rb.font.bold = True
    rb.font.color.rgb = RGBColor(0, 180, 150)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.line_spacing = 1.1
    rt = p_title.add_run("LeadCraft AI")
    rt.font.name = "Calibri"
    rt.font.size = Pt(28)
    rt.font.bold = True
    rt.font.color.rgb = RGBColor(10, 25, 47)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    p_sub.paragraph_format.line_spacing = 1.2
    rs = p_sub.add_run("Servicio Integral de Gestión y Presencia Digital para el Pequeño Comercio Local impulsado por Inteligencia Artificial")
    rs.font.name = "Calibri"
    rs.font.size = Pt(12)
    rs.font.italic = True
    rs.font.color.rgb = RGBColor(0, 56, 255)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(10)
    rd = p_div.add_run("—" * 32)
    rd.font.color.rgb = RGBColor(203, 213, 225)

    p_hito = doc.add_paragraph()
    p_hito.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hito.paragraph_format.space_after = Pt(25)
    p_hito.paragraph_format.line_spacing = 1.25
    rh = p_hito.add_run("HITO 1: ANÁLISIS Y CONTEXTO\nDocumento de Viabilidad Inicial, Definición de Alcance y Modelo de Servicio")
    rh.font.name = "Calibri"
    rh.font.size = Pt(11)
    rh.font.bold = True
    rh.font.color.rgb = RGBColor(71, 85, 105)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.line_spacing = 1.3
    p_meta.paragraph_format.space_after = Pt(28)
    
    p_meta.add_run("Alumno: ").font.bold = True
    p_meta.add_run("Kevin Muñoz Córcoles\n")
    p_meta.add_run("Curso: ").font.bold = True
    p_meta.add_run("2º Desarrollo de Aplicaciones Multiplataforma (DAM)\n")
    p_meta.add_run("Módulo: ").font.bold = True
    p_meta.add_run("Proyecto Intermodular | Convocatoria 2026/2027\n")
    p_meta.add_run("Fecha de entrega: ").font.bold = True
    p_meta.add_run("23 de Octubre de 2026")

    add_h1("Índice del Documento de Viabilidad Inicial", space_before=6, space_after=3)
    toc_lines = [
        "01. Título y Descripción del Proyecto: El Servicio LeadCraft AI (Pág. 2)",
        "02. Problema o Necesidad: La Dificultad de Llevar un Negocio en Internet (Pág. 2)",
        "03. Usuarios y Público Objetivo (Pág. 2)",
        "04. Benchmarking Competitivo y Soluciones Similares (Pág. 2)",
        "05. Objetivos Iniciales del Proyecto (Modelo SMART) (Pág. 3)",
        "06. Alcance Mínimo Viable (MVP) y Ampliaciones Futuras (Pág. 3)",
        "07. Riesgos Iniciales y Medidas Preventivas (Pág. 3)",
        "08. Vinculación con los Objetivos de Desarrollo Sostenible (ODS) (Pág. 4)",
        "09. Organización Inicial del Trabajo y Metodología (Pág. 4)",
        "10. Caso de Uso Principal: Prospección, Venta e Implantación del Servicio (Pág. 4)",
        "11. Diagrama de Arquitectura del Sistema (Pág. 5)",
        "12. Mockups de Pantallas Clave de la Aplicación (Pág. 5)"
    ]
    for tl in toc_lines:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.line_spacing = 1.15
        p_t.paragraph_format.space_after = Pt(1.5)
        r_t = p_t.add_run("• " + tl)
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_page_break()

    # ========================== PÁGINA 2: ANÁLISIS & BENCHMARKING ==========================
    add_h1("01. Título y Descripción del Proyecto: El Servicio LeadCraft AI")
    add_p("LeadCraft AI es una solución de software como servicio (SaaS) diseñada para comercializar un servicio integral de digitalización y gestión online para pequeños comercios y profesionales locales (barberías, talleres, clínicas de fisioterapia, restaurantes o tiendas de barrio). El objetivo principal del proyecto no es únicamente crear páginas web estáticas, sino vender un servicio activo que ayuda al comerciante a llevar su negocio en Internet en el día a día: incluye la creación y alojamiento de su página web profesional mediante IA, un sistema integrado para gestionar sus citas y reservas online, catálogo de servicios actualizado, contacto directo por WhatsApp con sus clientes y un panel de control intuitivo desde el móvil o tablet para consultar sus reservas diarias y horarios sin depender de informáticos externos ni costes abusivos.")

    add_h1("02. Problema o Necesidad: La Dificultad de Llevar un Negocio en Internet")
    add_p("Hoy en día, la mayoría de los pequeños comerciantes y autónomos se encuentran desbordados para gestionar su presencia en Internet. Saben que necesitan estar digitalizados, pero no tienen tiempo, conocimientos técnicos ni presupuesto para contratar agencias de marketing o desarrolladores externos (que suelen cobrar entre 800€ y 2.500€ iniciales más cuotas de mantenimiento mensual desproporcionadas). Como consecuencia, muchos negocios solo tienen una ficha básica en Google Maps pero carecen de web propia, no ofrecen citas online y dependen de atender llamadas telefónicas mientras trabajan con un cliente en el local, perdiendo ventas frente a grandes franquicias. LeadCraft AI resuelve esta situación ofreciendo un servicio 'todo en uno' asequible y empaquetado que se encarga de poner en marcha y llevar su tienda o negocio en Internet de forma transparente.")

    add_h1("03. Usuarios y Público Objetivo")
    add_p("• Cliente Principal (El Comerciante / Dueño del Negocio): Propietarios de pymes locales que contratan el servicio de LeadCraft AI para despreocuparse de la complejidad técnica y tener su negocio operativo en Internet (web con citas, carta/catálogo y canal directo de WhatsApp).")
    add_p("• Clientes Finales del Negocio: Vecinos y usuarios de la zona que acceden a la web desde su móvil para reservar cita en la peluquería, encargar un pedido o consultar precios a cualquier hora.")
    add_p("• Operador Comercial / Gestor de Cuentas: Diseñadores freelance, consultores o alumnos de DAM que utilizan la herramienta sobre Google Maps para identificar comercios con carencias digitales en su zona y ofrecerles el servicio con una demo interactiva ya configurada.")

    add_h1("04. Benchmarking Competitivo y Soluciones Similares")
    add_p("Se han analizado tres alternativas actuales en el mercado para comparar cómo abordan la digitalización de pymes:")

    t_bench = doc.add_table(rows=1, cols=4)
    t_bench.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_bench.autofit = False

    b_hdrs = ["Plataforma", "Modelo / Enfoque", "Puntos Fuertes", "Carencias frente al Servicio LeadCraft"]
    b_widths = [Inches(1.4), Inches(1.5), Inches(1.8), Inches(2.0)]
    
    for i, h in enumerate(t_bench.rows[0].cells):
        h.text = b_hdrs[i]
        h.width = b_widths[i]
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A192F"/>')
        h._tc.get_or_add_tcPr().append(shd)
        for p in h.paragraphs:
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(255, 255, 255)

    bench_data = [
        ("Wix Studio / Duda", "Generadores web SaaS con IA para crear sitios web corporativos.", "Diseños modernos y plantillas responsive bien estructuradas.", "No incluye prospección en Maps ni detecta carencias; requiere que el comerciante dedique tiempo a construir y maquetar la web por su cuenta."),
        ("D7 Lead Finder", "Buscadores de leads que extraen datos de contacto desde Google Maps.", "Rápida extracción masiva de correos y teléfonos de negocios.", "Se limita a generar listados de texto en Excel; no audita carencias reales, no genera la solución web ni integra reservas."),
        ("Calendly / Treatwell", "Software especializado en reservas y gestión de citas online.", "Calendarios eficaces con recordatorios automáticos por SMS/correo.", "Son herramientas aisladas que no proporcionan la web del negocio ni canal directo por WhatsApp; comisiones por reserva abusivas en el caso de Treatwell.")
    ]

    for r_idx, (p_name, p_mod, p_str, p_car) in enumerate(bench_data):
        row = t_bench.add_row()
        fill = "F1F5F9" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate([p_name, p_mod, p_str, p_car]):
            cell = row.cells[c_idx]
            cell.width = b_widths[c_idx]
            cell.text = val
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>')
            cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(30, 41, 59)
                    if c_idx == 0:
                        r.font.bold = True

    b_tblPr = t_bench._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="0A192F"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    b_tblPr.append(borders)

    doc.add_page_break()

    # ========================== PÁGINA 3: SMART, MVP & RIESGOS ==========================
    add_h1("05. Objetivos Iniciales del Proyecto (Modelo SMART)")
    add_p("• Específico (S): Desarrollar una solución integral que permita activar un servicio digital completo para pequeños negocios locales (web profesional generada con IA mediante Google Gemini + widget de reservas online + panel de control simplificado para el comerciante).")
    add_p("• Medible (M): Conseguir que el tiempo de configuración y despliegue del servicio sea inferior a 60 segundos por comercio, y estructurar una oferta de servicio asequible para pymes con un coste inferior a 25€/mes frente a los costes tradicionales de agencia.")
    add_p("• Alcanzable (A): Desarrollar el sistema utilizando tecnologías multiplataforma dominadas en el ciclo (FastAPI en backend, Flutter/Web en frontend y bases de datos relacionales SQLite/PostgreSQL) con plantillas modulares prediseñadas por sector.")
    add_p("• Relevante (R): Cubre una necesidad real del mercado: transformar el tejido comercial de barrio dotándolo de autonomía digital y creando una oportunidad de autoempleo para titulados de DAM.")
    add_p("• Temporal (T): Cumplir la planificación académica en tres entregas: Hito 1 en octubre, Hito 2 (MVP funcional) en enero y Hito 3 (plataforma completa con panel de comerciante) en marzo/abril de 2027.")

    add_h1("06. Alcance Mínimo Viable (MVP) y Ampliaciones Futuras")
    add_p("• Funcionalidades del MVP (Versión 1.0 del Servicio):")
    mvp_items = [
        "1) Módulo de localización y auditoría express de comercios sobre Google Maps (detección de carencias web y reservas).",
        "2) Motor de generación con IA de la web del negocio (redacción comercial de servicios, horarios y fotos de Maps).",
        "3) Widget interactivo de citas y reservas online incrustado en la web con selección de franjas horarias.",
        "4) Panel de gestión básico para el comerciante donde puede consultar las citas recibidas y cambiar horarios desde su móvil.",
        "5) Integración de botón de contacto directo por WhatsApp Business.",
        "6) Exportación de la web en paquete HTML5/Tailwind y generación de informe comercial de auditoría en PDF."
    ]
    for mi in mvp_items:
        p_mi = doc.add_paragraph()
        p_mi.paragraph_format.line_spacing = 1.15
        p_mi.paragraph_format.space_after = Pt(1.5)
        p_mi.paragraph_format.left_indent = Inches(0.2)
        r_mi = p_mi.add_run(mi)
        r_mi.font.name = "Calibri"
        r_mi.font.size = Pt(9.5)
        r_mi.font.color.rgb = RGBColor(40, 50, 60)

    add_p("• Ampliaciones Futuras (Versión 2.0):")
    fut_items = [
        "1) Bot inteligente de WhatsApp con IA que responde preguntas de los clientes y agenda citas de forma automática 24/7.",
        "2) Pasarela de cobro integrada con Stripe para permitir a los clientes pagar señales o productos online.",
        "3) Alojamiento cloud automatizado en Vercel/Netlify con asignación de dominios propios (.es / .com) con 1 clic."
    ]
    for fi in fut_items:
        p_fi = doc.add_paragraph()
        p_fi.paragraph_format.line_spacing = 1.15
        p_fi.paragraph_format.space_after = Pt(1.5)
        p_fi.paragraph_format.left_indent = Inches(0.2)
        r_fi = p_fi.add_run(fi)
        r_fi.font.name = "Calibri"
        r_fi.font.size = Pt(9.5)
        r_fi.font.color.rgb = RGBColor(40, 50, 60)

    add_h1("07. Riesgos Iniciales y Medidas Preventivas")
    add_p("• Riesgo 1: Coste de las llamadas a Google Places API y modelos de IA. Medida: Implementar almacenamiento en caché local (base de datos) para no repetir consultas sobre el mismo comercio y comprimir los prompts enviados a Gemini.")
    add_p("• Riesgo 2: Reticencia del comerciante a la tecnología o falta de uso. Medida: Diseñar el panel de gestión con una interfaz ultrasimplificada tipo 'Modo Fácil', similar a WhatsApp, sin menús complicados ni términos técnicos.")
    add_p("• Riesgo 3: Gestión del tiempo durante el curso escolar. Medida: Priorización estricta de tareas según el método MoSCoW, asegurando que el flujo de venta y el panel de citas funcionen antes de añadir florituras estéticas.")

    doc.add_page_break()

    # ========================== PÁGINA 4: ODS, METODOLOGÍA & CASO DE USO ==========================
    add_h1("08. Vinculación con los Objetivos de Desarrollo Sostenible (ODS)")
    add_p("• ODS 8: Trabajo Decente y Crecimiento Económico (Meta 8.3): LeadCraft AI impulsa la competitividad y modernización de las microempresas y trabajadores autónomos de barrio, dándoles herramientas para incrementar sus ventas y competir con grandes cadenas multinacionales.")
    add_p("• ODS 9: Industria, Innovación e Infraestructura (Meta 9.3): Democratiza el acceso a la inteligencia artificial generativa y a infraestructuras web avanzadas para comercios tradicionales con presupuestos modestos.")
    add_p("• ODS 11: Ciudades y Comunidades Sostenibles (Meta 11.a): Fortalece el comercio de proximidad, garantizando que los servicios esenciales y tiendas de barrio sigan vivos y conectados con los vecinos en la era digital.")

    add_h1("09. Organización Inicial del Trabajo y Metodología")
    add_p("El proyecto se desarrollará siguiendo el marco ágil Scrum apoyado en un tablero Kanban de Trello para la gestión del Product Backlog. El desarrollo se organiza en tres Sprints principales: Sprint 1 (Explorador sobre Maps, API de Places y motor de auditoría); Sprint 2 (Generador web con IA de Gemini, catálogo de plantillas por sector y widget de citas online); Sprint 3 (Panel de control del comerciante, exportación cloud/PDF, integración de WhatsApp y pruebas de usuario).")

    add_h1("10. Caso de Uso Principal: Prospección, Venta e Implantación del Servicio")
    add_p("A continuación se detalla el flujo de funcionamiento y comercialización del servicio:")
    add_p("1. Detección en Maps: El gestor abre la aplicación, selecciona su zona geográfica y filtra por negocios de servicios (ej. barberías o clínicas dentales).")
    add_p("2. Diagnóstico de Carencias: El sistema analiza el local 'Barbería Clásica' y detecta que no tiene web ni reservas online (Score Digital del 45%).")
    add_p("3. Generación Inmediata de la Solución: La IA analiza los datos de Maps (horarios, opiniones y fotos) y genera en 25 segundos una web completa con precios de cortes y un widget de reservas de citas.")
    add_p("4. Presentación y Demostración: El comercial contacta con el dueño mostrando la web ya funcionando en el móvil: 'Hemos preparado esta demo para tu negocio; incluye tus citas y tu WhatsApp por una cuota reducida al mes'.")
    add_p("5. Activación del Servicio: El comerciante contrata el servicio; recibe acceso a su panel de control de LeadCraft AI en su móvil para gestionar sus citas entrantes y su web queda publicada en Internet.")

    doc.add_page_break()

    # ========================== PÁGINA 5: ARQUITECTURA & MOCKUPS ==========================
    add_h1("11. Diagrama de Arquitectura del Sistema", space_before=2, space_after=2)
    add_p("La arquitectura técnica de LeadCraft AI se estructura en cuatro capas modulares que separan el cliente, la lógica de negocio, las APIs externas y la persistencia de datos:", space_after=2)

    diag_path = os.path.join(curr_dir, "diagrama_arquitectura_leadcraft.png")
    if os.path.exists(diag_path):
        p_img_d = doc.add_paragraph()
        p_img_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img_d.paragraph_format.space_before = Pt(1)
        p_img_d.paragraph_format.space_after = Pt(1)
        r_d = p_img_d.add_run()
        r_d.add_picture(diag_path, width=Inches(5.5))
        
        p_cap_d = doc.add_paragraph()
        p_cap_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap_d.paragraph_format.space_after = Pt(4)
        rcd = p_cap_d.add_run("Figura 1: Arquitectura de Componentes de LeadCraft AI (Frontend, Backend Gateway, APIs Externas y Capa de Datos).")
        rcd.font.name = "Calibri"
        rcd.font.size = Pt(8)
        rcd.font.italic = True
        rcd.font.color.rgb = RGBColor(100, 116, 139)

    add_h1("12. Mockups de Pantallas Clave de la Aplicación", space_before=3, space_after=2)
    
    m1_path = os.path.join(curr_dir, "mockup_leadcraft_desktop.png")
    m2_path = os.path.join(curr_dir, "mockup_leadcraft_mobile.png")

    if os.path.exists(m1_path) and os.path.exists(m2_path):
        t_mock = doc.add_table(rows=1, cols=2)
        t_mock.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_mock.autofit = False

        c1, c2 = t_mock.rows[0].cells
        c1.width = Inches(3.35)
        c2.width = Inches(3.35)

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run()
        r1.add_picture(m1_path, width=Inches(3.15))
        p1_cap = c1.add_paragraph()
        p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1_cap.paragraph_format.space_before = Pt(1)
        p1_cap.paragraph_format.space_after = Pt(1)
        rc1 = p1_cap.add_run("Pantalla 1: Mapa y Auditoría Digital Express versión Escritorio.")
        rc1.font.size = Pt(8)
        rc1.font.bold = True
        rc1.font.color.rgb = RGBColor(71, 85, 105)

        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run()
        r2.add_picture(m2_path, width=Inches(3.15))
        p2_cap = c2.add_paragraph()
        p2_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2_cap.paragraph_format.space_before = Pt(1)
        p2_cap.paragraph_format.space_after = Pt(1)
        rc2 = p2_cap.add_run("Pantalla 2: Versión móvil en 3 fases (Explorar, Generar y Exportar).")
        rc2.font.size = Pt(8)
        rc2.font.bold = True
        rc2.font.color.rgb = RGBColor(71, 85, 105)

    # Save to BOTH filenames
    out_docx_leadcraft = os.path.join(curr_dir, "Hito1_LeadCraft_Kevin.docx")
    out_docx_viabilidad = os.path.join(curr_dir, "Hito1_Viabilidad_LeadCraft_Kevin.docx")
    doc.save(out_docx_leadcraft)
    doc.save(out_docx_viabilidad)
    print("Hito 1 Word Document saved at:", out_docx_leadcraft)
    print("Hito 1 Word Document saved at:", out_docx_viabilidad)

if __name__ == "__main__":
    build_word_hito1()
